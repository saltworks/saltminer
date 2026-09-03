import os

import pytest
from docx import Document
from docx.oxml.ns import qn

from smreport.convert import convert_file, field_expression
from tests.conftest import SHIPPED_TEMPLATE, paragraph_texts


def test_field_expression_attribute_mapping():
    assert field_expression("EngagementAttributes|tester", []) == "EngagementAttribute_tester"
    assert field_expression("IssueAttributes|tested_by", [("IssueDetails", "detail")]) == "detail.IssueAttribute_tested_by"
    assert field_expression("Name", [("AssetTocs", "asset"), ("Issues", "issue")]) == "issue.Name"


def test_convert_synthetic(mergefield_template, tmp_path):
    dest = str(tmp_path / "out.docx")
    report = convert_file(mergefield_template, dest)
    tags = [c.tag for c in report]
    assert "{%p for asset in AssetTocs %}" in tags
    assert "{%p for group in asset.SeverityGroups %}" in tags
    assert "{%p for issue in group.Issues %}" in tags
    assert "{%p for toc in IssueTocs %}" in tags
    assert "{%p for detail in IssueDetails %}" in tags
    assert "{{ detail.IssueAttribute_tested_by }}" in tags
    assert "{{ EngagementAttribute_tester }}" in tags
    texts = paragraph_texts(dest)
    assert not any("TableStart" in t or "«" in t for t in texts)
    assert "{{ Name }} - {{ State }}" in texts
    # bold formatting on the Name field survived conversion
    doc = Document(dest)
    p = next(p for p in doc.element.body.iter(qn("w:p")) if "{{ Name }}" in "".join(t.text for t in p.iter(qn("w:t"))))
    r = next(r for r in p.findall(qn("w:r")) if "{{ Name }}" in "".join(t.text or "" for t in r.findall(qn("w:t"))))
    assert r.find(qn("w:rPr")) is not None and r.find(qn("w:rPr")).find(qn("w:b")) is not None
    # no complex field remnants
    assert doc.element.body.find(f".//{qn('w:fldChar')}") is None


@pytest.mark.skipif(not os.path.exists(SHIPPED_TEMPLATE), reason="shipped template not present")
def test_convert_shipped_template(tmp_path):
    dest = str(tmp_path / "SaltworksTemplate-jinja.docx")
    report = convert_file(SHIPPED_TEMPLATE, dest)
    assert len(report) == 57
    assert sum(c.placement == "removed" for c in report) == 2       # Section1 start/end
    assert sum(c.tag.startswith("{%p for") for c in report) == 6
    assert sum(c.tag == "{%p endfor %}" for c in report) == 6
    doc = Document(dest)
    assert doc.element.body.find(f".//{qn('w:fldChar')}") is None
    assert doc.element.body.find(f".//{qn('w:tbl')}") is not None      # issue table preserved
    assert len(doc.sections[0].footer.paragraphs) > 0                     # footer preserved
