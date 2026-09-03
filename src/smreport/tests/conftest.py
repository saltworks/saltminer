import base64
import json
import os
import shutil
import struct
import zlib

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures")
REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
SHIPPED_TEMPLATE = os.path.join(
    REPO_ROOT, "Saltworks.SaltMiner.JobManager", "Saltworks.SaltMiner.JobManager",
    "TemplateDefaults", "Saltworks", "SaltworksTemplate.docx",
)


def make_png(width: int, height: int, rgb=(200, 30, 30)) -> bytes:
    """Minimal solid-colour PNG without Pillow."""
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(tag, data):
        c = tag + data
        return struct.pack(">I", len(data)) + c + struct.pack(">I", zlib.crc32(c) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b"")


def data_uri(png: bytes) -> str:
    return "data:image/png;base64," + base64.b64encode(png).decode()


@pytest.fixture
def big_png(tmp_path):
    p = tmp_path / "images" / "big.png"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(make_png(600, 200))   # 600px @72dpi = 8.33in wide -> must shrink to 3in
    return str(p)


@pytest.fixture
def context_file(tmp_path, big_png):
    """Copy of the hand-built fixture context with a data-URI image injected."""
    with open(os.path.join(FIXTURES, "context.json"), encoding="utf8") as fh:
        ctx = json.load(fh)
    ctx["IssueDetails"][0]["Details"] = (
        ctx["IssueDetails"][0]["Details"] + "\n\nInline data image: ![tiny](" + data_uri(make_png(40, 20)) + ")"
    )
    dest = tmp_path / "context.json"
    dest.write_text(json.dumps(ctx), encoding="utf8")
    return str(dest)


def add_mergefield(paragraph, name: str, bold=False):
    """Append a Word complex MERGEFIELD (begin/instr/separate/result/end runs)."""
    def run(child):
        r = OxmlElement("w:r")
        if bold:
            rpr = OxmlElement("w:rPr")
            rpr.append(OxmlElement("w:b"))
            r.append(rpr)
        r.append(child)
        paragraph._p.append(r)

    for kind in ("begin",):
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); run(fc)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" MERGEFIELD  {name} "; run(it)
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); run(fc)
    t = OxmlElement("w:t"); t.text = f"«{name}»"; run(t)
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); run(fc)


def build_mergefield_template(path: str, static_png: str | None = None):
    """A synthetic Syncfusion-style template covering every group in the DTO."""
    doc = Document()
    doc.add_paragraph()  # blank
    add_mergefield(doc.add_paragraph(), "TableStart:Section1")
    p = doc.add_paragraph(); add_mergefield(p, "Name", bold=True); p.add_run(" - "); add_mergefield(p, "State")
    p = doc.add_paragraph(); p.add_run("Created For: "); add_mergefield(p, "Customer")
    add_mergefield(doc.add_paragraph(), "EngagementAttributes|tester")
    add_mergefield(doc.add_paragraph(), "Summary")
    if static_png:
        doc.add_picture(static_png, width=Pt(400))
        docpr = doc.paragraphs[-1]._p.find(f".//{qn('wp:docPr')}")
        docpr.set("descr", "StaticImage")
    add_mergefield(doc.add_paragraph(), "TableStart:AssetTocs")
    p = doc.add_paragraph(); add_mergefield(p, "Name"); p.add_run(" - "); add_mergefield(p, "Id")
    add_mergefield(doc.add_paragraph(), "TableStart:SeverityGroups")
    add_mergefield(doc.add_paragraph(), "Severity")
    add_mergefield(doc.add_paragraph(), "TableStart:Issues")
    p = doc.add_paragraph(); add_mergefield(p, "Name"); p.add_run(" Total: "); add_mergefield(p, "Total")
    add_mergefield(doc.add_paragraph(), "TableEnd:Issues")
    add_mergefield(doc.add_paragraph(), "TableEnd:SeverityGroups")
    add_mergefield(doc.add_paragraph(), "TableEnd:AssetTocs")
    add_mergefield(doc.add_paragraph(), "TableStart:IssueTocs")
    add_mergefield(doc.add_paragraph(), "Severity")
    add_mergefield(doc.add_paragraph(), "TableStart:SeverityGroups")
    p = doc.add_paragraph(); add_mergefield(p, "Name"); p.add_run(" Total: "); add_mergefield(p, "Total")
    add_mergefield(doc.add_paragraph(), "TableEnd:SeverityGroups")
    add_mergefield(doc.add_paragraph(), "TableEnd:IssueTocs")
    add_mergefield(doc.add_paragraph(), "TableStart:IssueDetails")
    add_mergefield(doc.add_paragraph(), "Name")
    table = doc.add_table(rows=7, cols=2)
    for row, name in zip(table.rows, ("Severity", "Proof", "Details", "Implication", "Recommendation", "References", "IssueAttributes|tested_by")):
        row.cells[0].text = name
        add_mergefield(row.cells[1].paragraphs[0], name)
    p = doc.add_paragraph(); p.add_run("Comments: "); add_mergefield(p, "Comments")
    add_mergefield(doc.add_paragraph(), "TableEnd:IssueDetails")
    add_mergefield(doc.add_paragraph(), "TableEnd:Section1")
    doc.save(path)
    return path


@pytest.fixture
def mergefield_template(tmp_path, big_png):
    return build_mergefield_template(str(tmp_path / "merge.docx"), static_png=big_png)


@pytest.fixture
def jinja_template(tmp_path, mergefield_template):
    from smreport.convert import convert_file
    dest = str(tmp_path / "merge-jinja.docx")
    convert_file(mergefield_template, dest)
    return dest


def paragraph_texts(path: str) -> list[str]:
    doc = Document(path)
    out = []
    for p in doc.element.body.iter(qn("w:p")):
        out.append("".join(t.text or "" for t in p.iter(qn("w:t"))))
    return out


def soffice_available() -> bool:
    return shutil.which("soffice") is not None
