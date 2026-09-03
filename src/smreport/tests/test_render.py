import os

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from smreport.render import RenderError, RenderOptions, render
from tests.conftest import paragraph_texts


def _render(context_file, template, tmp_path, **kw):
    outdir = str(tmp_path / "out")
    opts = RenderOptions(context_path=context_file, template_path=template, outdir=outdir, name="report", **kw)
    return render(opts)


def _runs(p):
    return [r for r in p.iter(qn("w:r"))]


def _run_text(r):
    return "".join(t.text or "" for t in r.findall(qn("w:t")))


def test_end_to_end_docx(context_file, jinja_template, tmp_path):
    out = _render(context_file, jinja_template, tmp_path, value_colors={"critical": "Red"})
    assert set(out) == {"docx"} and os.path.isfile(out["docx"])
    texts = paragraph_texts(out["docx"])
    joined = "\n".join(texts)

    # every group rendered
    assert "Acme Web Portal Pentest - Active" in texts
    assert "portal.acme.example - asset-1" in texts and "api.acme.example - asset-2" in texts
    assert "SQL Injection Total: 1" in texts and "Verbose Errors Total: 1" in texts
    assert texts.count("SQL Injection Total: 1") == 2      # once under AssetTocs, once under IssueTocs
    assert "Cameron" in texts                              # EngagementAttributes|tester mapped to DTO name
    assert "cameron" in texts and "alex" in texts          # IssueAttributes|tested_by per issue
    assert joined.count("Verbose Errors") >= 3             # toc + detail
    # no Jinja/sentinel residue, no "None"
    assert "{{" not in joined and "{%" not in joined and "" not in joined
    assert "None" not in texts
    # summary newline became a line break inside one paragraph
    p = next(p for p in Document(out["docx"]).element.body.iter(qn("w:p")) if "Second line of summary" in "".join(t.text or "" for t in p.iter(qn("w:t"))))
    assert p.find(f".//{qn('w:br')}") is not None


def test_markdown_rich_text(context_file, jinja_template, tmp_path):
    out = _render(context_file, jinja_template, tmp_path)
    doc = Document(out["docx"])
    body = doc.element.body
    paras = list(body.iter(qn("w:p")))

    def find(text):
        return next(p for p in paras if text in "".join(t.text or "" for t in p.iter(qn("w:t"))))

    # bold run inside markdown
    p = find("endpoint concatenates")
    bold_runs = [r for r in _runs(p) if r.find(qn("w:rPr")) is not None and r.find(qn("w:rPr")).find(qn("w:b")) is not None]
    assert any(_run_text(r) == "login" for r in bold_runs)
    # list items got numbering
    item = find("Parameter: ")
    assert item.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None
    nested = find("includes ")
    assert nested.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:ilvl')}").get(qn("w:val")) == "1"
    ordered = find("first")
    num_id = ordered.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}").get(qn("w:val"))
    assert num_id != item.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}").get(qn("w:val"))
    # code block kept its lines and mono font
    code = find("POST /login")
    assert code.find(f".//{qn('w:rFonts')}").get(qn("w:ascii")) == "Consolas"
    assert code.find(f".//{qn('w:br')}") is not None
    # table from markdown
    assert any("bio" in "".join(t.text or "" for t in tbl.iter(qn("w:t"))) for tbl in body.iter(qn("w:tbl")))
    # heading + rule + quote
    assert find("Heading in markdown").find(f".//{qn('w:b')}") is not None
    assert find("Quoted note").find(f"{qn('w:pPr')}/{qn('w:ind')}") is not None
    # strike
    p = find("customer records")
    assert any(r.find(qn("w:rPr")) is not None and r.find(qn("w:rPr")).find(qn("w:strike")) is not None and _run_text(r) == "some" for r in _runs(p))
    # broken image skipped, text preserved
    assert "and text continues." in "".join(t.text or "" for t in find("Broken image").iter(qn("w:t")))


def test_hyperlinks_blue_underlined(context_file, jinja_template, tmp_path):
    out = _render(context_file, jinja_template, tmp_path)
    doc = Document(out["docx"])
    links = list(doc.element.body.iter(qn("w:hyperlink")))
    assert links, "markdown links should produce w:hyperlink elements"
    for h in links:
        for r in h.findall(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            assert rpr.find(qn("w:color")).get(qn("w:val")) == "0000FF"
            assert rpr.find(qn("w:u")).get(qn("w:val")) == "single"
    rels = doc.part.rels
    targets = {rel.target_ref for rel in rels.values() if rel.is_external}
    assert "https://cheatsheetseries.owasp.org/" in targets
    assert "https://owasp.org/Top10/A03_2021-Injection/" in targets    # autolink


def test_images_resized_static_skipped(context_file, jinja_template, tmp_path):
    out = _render(context_file, jinja_template, tmp_path, image_max_width=216, image_max_height=288)
    doc = Document(out["docx"])
    inlines = list(doc.element.body.iter(qn("wp:inline")))
    assert len(inlines) == 3   # static template image + markdown big.png + data-URI tiny
    by_alt = {}
    for inl in inlines:
        docpr = inl.find(qn("wp:docPr"))
        ext = inl.find(qn("wp:extent"))
        by_alt[docpr.get("descr")] = (int(ext.get("cx")), int(ext.get("cy")))
    assert by_alt["StaticImage"][0] == int(Pt(400))             # untouched
    assert by_alt["login bypass"][0] == int(Pt(216))            # width-bound resize
    assert abs(by_alt["login bypass"][1] - int(Pt(216) / 3)) <= 12700
    assert by_alt["tiny"][0] < int(Pt(216))                     # small image untouched


def test_value_colors(context_file, jinja_template, tmp_path):
    out = _render(context_file, jinja_template, tmp_path, value_colors={"critical": "Red"})
    doc = Document(out["docx"])
    colored = [r for r in doc.element.body.iter(qn("w:r"))
               if _run_text(r) == "Critical" and r.find(qn("w:rPr")) is not None
               and r.find(qn("w:rPr")).find(qn("w:color")) is not None
               and r.find(qn("w:rPr")).find(qn("w:color")).get(qn("w:val")) == "FF0000"]
    assert colored


def test_missing_inputs_and_bad_format(context_file, jinja_template, tmp_path):
    with pytest.raises(RenderError, match="not found"):
        render(RenderOptions(context_path="/nope.json", template_path=jinja_template, outdir=str(tmp_path)))
    with pytest.raises(RenderError, match="unsupported format"):
        render(RenderOptions(context_path=context_file, template_path=jinja_template, outdir=str(tmp_path), formats=("xlsx",)))


def test_cli_exit_codes(context_file, jinja_template, tmp_path):
    from smreport.cli import main
    assert main(["--context", context_file, "--template", jinja_template, "--outdir", str(tmp_path / "o"), "--name", "r"]) == 0
    assert os.path.isfile(tmp_path / "o" / "r.docx")
    assert main(["--context", "/nope.json", "--template", jinja_template, "--outdir", str(tmp_path)]) == 2
    assert main(["--context", context_file, "--template", jinja_template, "--outdir", str(tmp_path), "--formats", "pdf", "--soffice", "/definitely/not/soffice"]) == 4
