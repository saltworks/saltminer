''' --[auto-generated, do not modify this block]--
*
* SaltMiner - The open source vulnerability and pen testing management platform
* Copyright (C) 2024-2026 Saltworks Security, LLC
*
* This program is free software: you can redistribute it and/or modify
* it under the terms of the GNU General Public License as published by
* the Free Software Foundation, either version 3 of the License.
*
* This program is distributed in the hope that it will be useful,
* but WITHOUT ANY WARRANTY; without even the implied warranty of
* MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the
* GNU General Public License for more details.
*
* You should have received a copy of the GNU General Public License
* along with this program. If not, see <https://www.gnu.org/licenses/>.
*
* ----
'''
# PBI-083: markdown images resolve, embed, and fall back, against a recording stub transport.

import base64
import io
import logging
import unittest
import zipfile
from unittest import mock

import docx
from docx.oxml.ns import qn

from Reports.Images import (ImageEmbedder, ImageResolver, ImageResolveError, UiApiFileClient,
                            http_transport)
from Reports.Markdown import render_markdown
from Reports.Tests.synthetic import add_field, paragraph_text, png_bytes
from Reports.ValueRenderer import make_value_renderer
from Reports.WordMerge import bind_roots, merge_document

GUID = "0b8c6d1e-1111-4222-8333-444455556666"
BASE = "http://ui-api:5001"
KEY_HEADER = "ReportingAuthorization"
KEY = "secret-reporting-key"


class RecordingTransport:
    def __init__(self, responses=None, default=None):
        self.calls = []
        self._responses = responses or {}
        self._default = default if default is not None else (200, png_bytes(40, 30))

    def __call__(self, url, headers):
        self.calls.append((url, dict(headers)))
        return self._responses.get(url, self._default)


def _blank_document():
    document = docx.Document()
    for paragraph in list(document.paragraphs):
        paragraph._p.getparent().remove(paragraph._p)
    return document


def _embedder(transport, foreign=None, with_client=True):
    client = UiApiFileClient(BASE, KEY_HEADER, KEY, transport) if with_client else None
    return ImageEmbedder(ImageResolver(client, foreign=foreign or transport))


def _render(document, embedder, fields):
    """Merge one record whose `fields` (name -> markdown) are all markdown fields."""
    add_field(document.add_paragraph(), "TableStart:Section1")
    for name in fields:
        add_field(document.add_paragraph(), name)
    add_field(document.add_paragraph(), "TableEnd:Section1")
    embedder.bind(document)
    renderer, result = make_value_renderer(set(fields), {}, images=embedder)
    merge_document(document, bind_roots(dict(fields)), renderer=renderer)
    return result


def _save(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return zipfile.ZipFile(io.BytesIO(buffer.getvalue()))


def _drawings(document):
    return list(document.element.body.iter(qn("w:drawing")))


def _text(document):
    return "".join(paragraph_text(p) for p in document.element.body.iter(qn("w:p")))


class AppAndForeignImages(unittest.TestCase):
    def test_app_url_is_rewritten_with_key_and_foreign_url_is_fetched_unchanged_without_it(self):
        # AC-1
        transport = RecordingTransport()
        document = _blank_document()
        proof = (f"![shot](https://example.test/api/File/{GUID}.png)\n\n"
                 "![ext](https://external.test/img/a.png)")
        _render(document, _embedder(transport), {"Proof": proof})

        self.assertEqual(transport.calls, [
            (f"{BASE}/File/{GUID}.png", {KEY_HEADER: KEY}),
            ("https://external.test/img/a.png", {}),
        ])
        self.assertEqual(len(_drawings(document)), 2)
        self.assertNotIn("[shot]", _text(document))
        self.assertNotIn("[ext]", _text(document))
        ids = [d.find(".//" + qn("wp:docPr")).get("id") for d in _drawings(document)]
        self.assertEqual(len(set(ids)), 2)

    def test_attachment_suffix_is_dropped_from_the_app_route(self):
        transport = RecordingTransport()
        _render(_blank_document(), _embedder(transport),
                {"Proof": f"![a](https://example.test/File/{GUID}/attachment)"})
        self.assertEqual(transport.calls, [(f"{BASE}/File/{GUID}", {KEY_HEADER: KEY})])

    def test_alt_text_becomes_the_picture_description(self):
        document = _blank_document()
        _render(document, _embedder(RecordingTransport()), {"Proof": "![my shot](https://x.test/a.png)"})
        descr = _drawings(document)[0].find(".//" + qn("wp:docPr")).get("descr")
        self.assertEqual(descr, "my shot")

    def test_repeated_url_is_fetched_once(self):
        transport = RecordingTransport()
        _render(_blank_document(), _embedder(transport),
                {"Proof": "![a](https://x.test/a.png)", "ProofImgs": "![a](https://x.test/a.png)"})
        self.assertEqual(len(transport.calls), 1)


class DataUriImage(unittest.TestCase):
    def test_data_uri_in_a_custom_markdown_attribute_embeds_identical_bytes(self):
        # AC-2
        png = png_bytes(10, 10)
        uri = "data:image/png;base64," + base64.b64encode(png).decode()
        transport = RecordingTransport()
        document = _blank_document()
        _render(document, _embedder(transport), {"IssueAttribute_Evidence": f"![d]({uri})"})

        self.assertEqual(len(_drawings(document)), 1)
        archive = _save(document)
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
        self.assertEqual(len(media), 1)
        self.assertEqual(archive.read(media[0]), png)
        self.assertEqual(transport.calls, [])


class FailureFallsBack(unittest.TestCase):
    def test_a_404_for_one_image_keeps_alt_text_and_warns(self):
        # AC-4
        bad = f"https://example.test/File/{GUID}.png"
        transport = RecordingTransport({f"{BASE}/File/{GUID}.png": (404, b"")})
        document = _blank_document()
        embedder = _embedder(transport)
        with self.assertLogs(level=logging.WARNING):
            result = _render(document, embedder,
                             {"Proof": f"![bad]({bad})\n\n![good](https://x.test/g.png)"})

        self.assertEqual(len(_drawings(document)), 1)
        self.assertIn("[bad]", _text(document))
        italic = [r for r in document.element.body.iter(qn("w:r"))
                  if paragraph_text(r) == "[bad]" and r.find(qn("w:rPr")).find(qn("w:i")) is not None]
        self.assertEqual(len(italic), 1)
        self.assertEqual(len(result.warnings), 1)
        self.assertIn(bad, result.warnings[0])
        self.assertIn("HTTP 404", result.warnings[0])

    def test_500_with_a_json_body_and_a_200_empty_body_both_fall_back(self):
        transport = RecordingTransport({
            f"{BASE}/File/{GUID}.png": (500, b'{"message":"x"}'),
            "https://x.test/empty.png": (200, b""),
        })
        document = _blank_document()
        with self.assertLogs(level=logging.WARNING):
            result = _render(document, _embedder(transport), {
                "Proof": f"![a](https://e.test/File/{GUID}.png)\n\n![b](https://x.test/empty.png)"})
        self.assertEqual(len(_drawings(document)), 0)
        self.assertEqual(len(result.warnings), 2)

    def test_a_transport_exception_falls_back(self):
        def boom(url, headers):
            raise ConnectionError("no route")
        document = _blank_document()
        with self.assertLogs(level=logging.WARNING):
            result = _render(document, _embedder(boom), {"Proof": "![a](https://x.test/a.png)"})
        self.assertEqual(len(_drawings(document)), 0)
        self.assertIn("ConnectionError", result.warnings[0])

    def test_unrecognised_image_bytes_fall_back(self):
        transport = RecordingTransport(default=(200, b"<svg></svg>"))
        document = _blank_document()
        with self.assertLogs(level=logging.WARNING):
            result = _render(document, _embedder(transport), {"Proof": "![a](https://x.test/a.svg)"})
        self.assertEqual(len(_drawings(document)), 0)
        self.assertIn("unrecognised image format", result.warnings[0])

    def test_the_warning_for_a_data_uri_is_truncated(self):
        uri = "data:image/png;base64," + "A" * 500
        embedder = _embedder(RecordingTransport())
        embedder.bind(_blank_document())
        with self.assertLogs(level=logging.WARNING):
            self.assertIsNone(embedder.run_for(uri, "x", None))
        self.assertLess(len(embedder.warnings[0]), 200)


class Resolver(unittest.TestCase):
    def test_foreign_url_is_fetched_with_no_client(self):
        transport = RecordingTransport()
        resolver = ImageResolver(None, foreign=transport)
        resolver.resolve("HTTPS://x.test/a.png")
        self.assertEqual(transport.calls, [("HTTPS://x.test/a.png", {})])

    def test_app_url_with_no_client_raises(self):
        resolver = ImageResolver(None, foreign=RecordingTransport())
        with self.assertRaises(ImageResolveError) as ctx:
            resolver.resolve(f"https://e.test/File/{GUID}.png")
        self.assertEqual(str(ctx.exception), "no UI API address")

    def test_other_schemes_and_relative_paths_raise_without_a_call(self):
        transport = RecordingTransport()
        resolver = ImageResolver(UiApiFileClient(BASE, KEY_HEADER, KEY, transport), foreign=transport)
        for src in ("ftp://x.test/a.png", f"/api/File/{GUID}.png", "file:///etc/passwd",
                    "data:text/plain;base64,AAAA"):
            with self.assertRaises(ImageResolveError):
                resolver.resolve(src)
        self.assertEqual(transport.calls, [])

    def test_trailing_slash_in_the_base_url_does_not_double(self):
        transport = RecordingTransport()
        UiApiFileClient(BASE + "/", KEY_HEADER, KEY, transport).fetch(GUID)
        self.assertEqual(transport.calls[0][0], f"{BASE}/File/{GUID}")


class HttpTransport(unittest.TestCase):
    def test_redirects_are_never_followed_for_keyed_or_keyless_calls(self):
        response = mock.Mock(status_code=200, content=b"x")
        with mock.patch("requests.Session.get", return_value=response) as get:
            transport = http_transport(False, 3)
            transport("http://a.test/File/x", {KEY_HEADER: KEY})
            transport("http://b.test/x.png", {})
        for call in get.call_args_list:
            self.assertIs(call.kwargs["allow_redirects"], False)
            self.assertEqual(call.kwargs["timeout"], 3)
            self.assertIs(call.kwargs["verify"], False)


class WithoutImages(unittest.TestCase):
    def test_render_markdown_without_images_still_prints_alt_text(self):
        paragraphs = render_markdown("![shot](https://x.test/a.png)")
        self.assertEqual(paragraph_text(paragraphs[0]), "[shot]")


if __name__ == "__main__":
    unittest.main()
