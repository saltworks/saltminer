''' --[auto-generated, do not modify this block]--
*
* SaltMiner - The open source vulnerability and pen testing management platform
* Copyright (C) 2024-2026 Saltworks Security, LLC
*
* This program is free software: you can redistribute it and/or modify
* it under the terms of the GNU General Public License as published by
* the Free Software Foundation, either version 3 of the License.
*
* This program is distributed in the hope that it will be useful,
* but WITHOUT ANY WARRANTY; without even the implied warranty of
* MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the
* GNU General Public License for more details.
*
* You should have received a copy of the GNU General Public License
* along with this program. If not, see <https://www.gnu.org/licenses/>.
*
* ----
'''
import unittest

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from Reports.WordFields import iter_merge_fields, make_run, parse_field_name, replace_field
from Reports.Tests.synthetic import add_field, paragraph_text


class ParseFieldNameTests(unittest.TestCase):
    def test_switches_and_quotes_are_not_part_of_the_name(self):
        self.assertEqual(parse_field_name(" MERGEFIELD  Severity \\* MERGEFORMAT "), "Severity")
        self.assertEqual(parse_field_name(' MERGEFIELD "Issue Attributes|tested_by" '), "Issue Attributes|tested_by")
        self.assertEqual(parse_field_name("MERGEFIELD Name"), "Name")


class IterMergeFieldsTests(unittest.TestCase):
    def test_nested_field_leaves_outer_field_intact(self):
        document = docx.Document()
        paragraph = document.add_paragraph()
        outer = add_field(paragraph, "Outer", result_text="«Outer»", inner_name="Inner")
        fields = iter_merge_fields(paragraph._p)
        self.assertEqual(sorted(f.name for f in fields), ["Inner", "Outer"])
        self.assertIsNotNone(outer)

    def test_replace_field_removes_all_five_runs(self):
        document = docx.Document()
        paragraph = document.add_paragraph()
        add_field(paragraph, "Name", result_text="«Name»")
        field = iter_merge_fields(paragraph._p)[0]
        self.assertEqual(len(field.all_runs()), 5)
        replace_field(field, [make_run("value", field.format_source)])
        self.assertEqual(paragraph_text(paragraph._p), "value")
        self.assertEqual(len(paragraph._p.findall(qn("w:r"))), 1)
        self.assertEqual(iter_merge_fields(paragraph._p), [])


class MakeRunTests(unittest.TestCase):
    def test_newline_becomes_break_and_rpr_is_copied(self):
        source = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "44")
        rpr.append(size)
        source.append(rpr)
        run = make_run("a\nb", source)
        self.assertEqual(run.find(qn("w:rPr")).find(qn("w:sz")).get(qn("w:val")), "44")
        self.assertEqual(len(run.findall(qn("w:br"))), 1)
        self.assertIsNot(run.find(qn("w:rPr")), rpr)


if __name__ == "__main__":
    unittest.main()
