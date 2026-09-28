''' --[auto-generated, do not modify this block]--
*
* SaltMiner - The open source vulnerability and pen testing management platform
* Copyright (C) 2024-2026 Saltworks Security, LLC
*
* This program is free software: you can redistribute it and/or modify
* it under the terms of the GNU General Public License as published by
* the Free Software Foundation, either version 3 of the License.
*
* This program is distributed in the hope that it will be useful,
* but WITHOUT ANY WARRANTY; without even the implied warranty of
* MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the
* GNU General Public License for more details.
*
* You should have received a copy of the GNU General Public License
* along with this program. If not, see <https://www.gnu.org/licenses/>.
*
* ----
'''
import os
import tempfile
import unittest
import zipfile
from collections import Counter

import docx
from docx.oxml.ns import qn

from Reports.Tests.synthetic import add_complex_field, add_field, paragraph_text
from Reports.WordMerge import bind_roots, merge_document


def _blank_document():
    document = docx.Document()
    for paragraph in list(document.paragraphs):
        paragraph._p.getparent().remove(paragraph._p)
    return document


def _build_three_sections_template():
    """Section1 and Section2 hold only a Name field. Section3 also holds a row-level `Rows`
    group whose row carries a Name field, so both a body-level and a row-level SectionN group are
    exercised. Two Word sections (a mid-document section break plus the final one)."""
    document = _blank_document()

    add_field(document.add_paragraph(), "TableStart:Section1")
    add_field(document.add_paragraph(), "Name")
    add_field(document.add_paragraph(), "TableEnd:Section1")

    document.add_section()

    add_field(document.add_paragraph(), "TableStart:Section2")
    add_field(document.add_paragraph(), "Name")
    add_field(document.add_paragraph(), "TableEnd:Section2")

    add_field(document.add_paragraph(), "TableStart:Section3")
    add_field(document.add_paragraph(), "Name")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).paragraphs[0].add_run("Label")
    table.cell(0, 1).paragraphs[0].add_run("Value")
    table.cell(0, 2).paragraphs[0].add_run("Name")
    add_field(table.cell(1, 0).paragraphs[0], "TableStart:Rows")
    add_field(table.cell(1, 0).paragraphs[0], "Label")
    add_field(table.cell(1, 1).paragraphs[0], "Value")
    add_field(table.cell(1, 2).paragraphs[0], "Name")
    add_field(table.cell(1, 2).paragraphs[0], "TableEnd:Rows")
    add_field(document.add_paragraph(), "TableEnd:Section3")

    record = {
        "Name": "ENG-NAME",
        "Rows": [
            {"Label": "L1", "Value": "V1"},
            {"Label": "L2", "Value": "V2"},
            {"Label": "L3", "Value": "V3"},
        ],
    }
    return document, record


class RowGroups(unittest.TestCase):
    def test_same_row_repeats(self):
        document = _blank_document()
        add_field(document.add_paragraph(), "TableStart:Section1")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).paragraphs[0].add_run("Label")
        table.cell(0, 1).paragraphs[0].add_run("Value")
        add_field(table.cell(1, 0).paragraphs[0], "TableStart:Rows")
        add_field(table.cell(1, 0).paragraphs[0], "Label")
        add_field(table.cell(1, 1).paragraphs[0], "Value")
        add_field(table.cell(1, 1).paragraphs[0], "TableEnd:Rows")
        add_field(document.add_paragraph(), "TableEnd:Section1")

        record = {"Rows": [
            {"Label": "L1", "Value": "V1"},
            {"Label": "L2", "Value": "V2"},
            {"Label": "L3", "Value": "V3"},
        ]}
        result = merge_document(document, bind_roots(record))

        table = document.tables[0]
        self.assertEqual(len(table.rows), 4)
        self.assertEqual(table.cell(0, 0).text, "Label")
        self.assertEqual(table.cell(0, 1).text, "Value")
        for i, (label, value) in enumerate([("L1", "V1"), ("L2", "V2"), ("L3", "V3")], start=1):
            self.assertEqual(table.cell(i, 0).text, label)
            self.assertEqual(table.cell(i, 1).text, value)
        self.assertEqual(result.groups["Rows"], 3)
        body_text = "".join(paragraph_text(p) for p in document.element.body.iter(qn("w:p")))
        self.assertNotIn("TableStart", body_text)
        self.assertNotIn("TableEnd", body_text)

    def test_multi_row_block_repeats(self):
        document = _blank_document()
        add_field(document.add_paragraph(), "TableStart:Section1")
        table = document.add_table(rows=4, cols=1)
        add_field(table.cell(0, 0).paragraphs[0], "TableStart:Rows")
        add_field(table.cell(0, 0).paragraphs[0], "Field0")
        add_field(table.cell(1, 0).paragraphs[0], "Field1")
        add_field(table.cell(2, 0).paragraphs[0], "Field2")
        add_field(table.cell(3, 0).paragraphs[0], "Field3")
        add_field(table.cell(3, 0).paragraphs[0], "TableEnd:Rows")
        add_field(document.add_paragraph(), "TableEnd:Section1")

        record = {"Rows": [
            {"Field0": "a0", "Field1": "a1", "Field2": "a2", "Field3": "a3"},
            {"Field0": "b0", "Field1": "b1", "Field2": "b2", "Field3": "b3"},
        ]}
        merge_document(document, bind_roots(record))

        table = document.tables[0]
        self.assertEqual(len(table.rows), 8)
        expected = ["a0", "a1", "a2", "a3", "b0", "b1", "b2", "b3"]
        for i, text in enumerate(expected):
            self.assertEqual(table.cell(i, 0).text, text)

    def test_zero_records_leaves_header(self):
        document = _blank_document()
        add_field(document.add_paragraph(), "TableStart:Section1")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).paragraphs[0].add_run("Label")
        table.cell(0, 1).paragraphs[0].add_run("Value")
        add_field(table.cell(1, 0).paragraphs[0], "TableStart:Rows")
        add_field(table.cell(1, 0).paragraphs[0], "Label")
        add_field(table.cell(1, 1).paragraphs[0], "Value")
        add_field(table.cell(1, 1).paragraphs[0], "TableEnd:Rows")
        add_field(document.add_paragraph(), "TableEnd:Section1")

        result = merge_document(document, bind_roots({"Rows": []}))

        table = document.tables[0]
        self.assertEqual(len(table.rows), 1)
        self.assertEqual(table.cell(0, 0).text, "Label")
        self.assertEqual(result.groups["Rows"], 0)


class SectionRoots(unittest.TestCase):
    def test_every_section_root_once(self):
        document, record = _build_three_sections_template()

        self.assertEqual(len(list(document.element.body.iter(qn("w:sectPr")))), 2)

        result = merge_document(document, bind_roots(record))

        self.assertEqual(result.groups["Section1"], 1)
        self.assertEqual(result.groups["Section2"], 1)
        self.assertEqual(result.groups["Section3"], 1)

        table = document.tables[0]
        self.assertEqual(len(table.rows), 4)
        for i in range(1, 4):
            self.assertEqual(table.cell(i, 2).text, "ENG-NAME")

        # 3 section-level Name paragraphs plus one per Rows record, via outward fallback.
        name_paragraphs = [p for p in document.element.body.iter(qn("w:p"))
                           if paragraph_text(p) == "ENG-NAME"]
        self.assertEqual(len(name_paragraphs), 6)

    def test_non_merge_fields_untouched(self):
        document, record = _build_three_sections_template()
        add_complex_field(document.add_paragraph(), "TOC", "1")
        add_complex_field(document.add_paragraph(), "PAGEREF _Toc1", "1")
        add_complex_field(document.add_paragraph(), 'DOCPROPERTY "Issues"', "3")
        footer = document.sections[0].footer
        add_complex_field(footer.paragraphs[0], "PAGE", "1")
        add_complex_field(footer.add_paragraph(), "NUMPAGES", "5")

        def instr_counts():
            texts = list(document.element.body.iter(qn("w:instrText")))
            for section in document.sections:
                texts.extend(section.footer.part.element.iter(qn("w:instrText")))
            counts = Counter()
            for instr in texts:
                text = (instr.text or "").strip()
                if text:
                    counts[text.split()[0]] += 1
            return counts

        before = instr_counts()
        merge_document(document, bind_roots(record))
        after = instr_counts()

        for name in ("TOC", "PAGEREF", "DOCPROPERTY", "PAGE", "NUMPAGES"):
            self.assertEqual(after[name], before[name], name)
            self.assertGreater(after[name], 0, name)

        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "o.docx")
            document.save(path)
            with zipfile.ZipFile(path) as z:
                xml = z.read("word/document.xml").decode("utf-8")
        self.assertEqual(xml.count("MERGEFIELD"), 0)


if __name__ == "__main__":
    unittest.main()
