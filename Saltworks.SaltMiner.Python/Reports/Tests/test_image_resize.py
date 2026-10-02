''' --[auto-generated, do not modify this block]--
*
* SaltMiner - The open source vulnerability and pen testing management platform
* Copyright (C) 2024-2026 Saltworks Security, LLC
*
* This program is free software: you can redistribute it and/or modify
* it under the terms of the GNU General Public License as published by
* the Free Software Foundation, either version 3 of the License.
*
* This program is distributed in the hope that it will be useful,
* but WITHOUT ANY WARRANTY; without even the implied warranty of
* MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the
* GNU General Public License for more details.
*
* You should have received a copy of the GNU General Public License
* along with this program. If not, see <https://www.gnu.org/licenses/>.
*
* ----
'''
# PBI-083: resize_pictures applies the .NET resize arithmetic (ReportProcessor.cs 294-335).

import io
import unittest

import docx
from docx.oxml.ns import qn

from Reports.Images import EMU_PER_POINT, resize_pictures
from Reports.Tests.synthetic import png_bytes


def _document_with_pictures():
    document = docx.Document()
    paragraph = document.add_paragraph()
    for width, height, descr in ((2000, 1000, "wide"), (500, 2000, "tall"), (2000, 1000, "StaticImage"),
                                 (100, 50, "small")):
        inline = document.part.new_pic_inline(io.BytesIO(png_bytes(width, height)))
        inline.find(qn("wp:docPr")).set("descr", descr)
        inline.find(qn("wp:docPr")).set("id", str(1000 + len(list(document.element.body.iter(qn("w:drawing"))))))
        run = paragraph.add_run()
        drawing = run._r.makeelement(qn("w:drawing"), {})
        drawing.append(inline)
        run._r.append(drawing)
    return document


def _extents(document):
    out = {}
    for inline in document.element.body.iter(qn("wp:inline")):
        extent = inline.find(qn("wp:extent"))
        out[inline.find(qn("wp:docPr")).get("descr")] = (int(extent.get("cx")), int(extent.get("cy")))
    return out


class Resize(unittest.TestCase):
    def test_caps_ratio_static_exemption_and_setting_override(self):
        # AC-3
        document = _document_with_pictures()
        native = _extents(document)
        self.assertEqual(native["wide"], (2000 * EMU_PER_POINT, 1000 * EMU_PER_POINT))

        count = resize_pictures(document)
        sizes = _extents(document)
        self.assertEqual(count, 2)
        self.assertAlmostEqual(sizes["wide"][0], 216 * EMU_PER_POINT, delta=1)
        self.assertAlmostEqual(sizes["wide"][1], 108 * EMU_PER_POINT, delta=1)
        self.assertAlmostEqual(sizes["tall"][1], 288 * EMU_PER_POINT, delta=1)
        self.assertAlmostEqual(sizes["tall"][0], 72 * EMU_PER_POINT, delta=1)
        self.assertEqual(sizes["StaticImage"], native["StaticImage"])
        self.assertEqual(sizes["small"], native["small"])

        narrow = _document_with_pictures()
        resize_pictures(narrow, max_width_pt=100)
        self.assertAlmostEqual(_extents(narrow)["wide"][0], 100 * EMU_PER_POINT, delta=1)

    def test_the_picture_transform_extent_follows_the_inline_extent(self):
        document = _document_with_pictures()
        resize_pictures(document)
        inline = next(document.element.body.iter(qn("wp:inline")))
        extent = inline.find(qn("wp:extent"))
        ext = next(inline.iter(qn("a:ext")))
        self.assertEqual((ext.get("cx"), ext.get("cy")), (extent.get("cx"), extent.get("cy")))

    def test_static_alt_text_is_read_from_the_setting(self):
        document = _document_with_pictures()
        resize_pictures(document, static_alt="wide")
        self.assertEqual(_extents(document)["wide"], (2000 * EMU_PER_POINT, 1000 * EMU_PER_POINT))


if __name__ == "__main__":
    unittest.main()
